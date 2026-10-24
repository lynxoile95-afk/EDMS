"""MDL ingestion from Excel with a full validation report.

The import is **all-or-nothing**: if any row has a blocking error the whole
batch is rejected and nothing is written.  Warnings are allowed and recorded.
Column mapping is explicit (never guessed silently) but suggested from header
aliases in English and Persian.
"""

from __future__ import annotations

import hashlib
import io
import re
import uuid
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Discipline, MDLImportBatch, MDLItem, Project, Unit, User
from app.models.enums import MDLImportStatus, Severity
from app.services import audit
from app.services.normalization import (
    is_valid_code,
    normalize_code,
    normalize_rev,
    normalize_title,
)
from app.services.project_settings import merge_settings
from app.services.revisions import rev_to_seq
from app.utils.jalali import parse_date

#: standard field -> accepted header aliases (normalised: lower, no spaces/dots)
COLUMN_ALIASES: dict[str, list[str]] = {
    "doc_code": ["documentno", "documentnumber", "docno", "doccode", "documentcode", "drawingno",
                 "drawingnumber", "dwgno", "code", "کدمدرک", "کدنقشه", "شمارهمدرک", "شمارهنقشه"],
    "title": ["title", "documenttitle", "description", "docname", "subject", "عنوان", "عنوانمدرک",
              "شرح", "موضوع", "ناممدرک"],
    "discipline": ["discipline", "dept", "department", "disciplinecode", "رشته", "دیسیپلین", "دپارتمان"],
    "unit": ["unit", "unitcode", "area", "یونیت", "واحد"],
    "doc_type": ["type", "documenttype", "doctype", "نوع", "نوعمدرک"],
    "current_rev": ["rev", "revision", "currentrev", "currentrevision", "بازنگری", "ویرایش", "ریویژن"],
    "status": ["status", "docstatus", "وضعیت", "وضعیتمدرک"],
    "responsible": ["responsible", "owner", "assignee", "resp", "مسئول", "مسئولتهیه"],
    "due_date": ["duedate", "planneddate", "targetdate", "موعد", "تاریختحویل", "تاریخموعد"],
    "planned_issue_date": ["issuedate", "plannedissuedate", "تاریخصدور", "تاریخبرنامه"],
    "weight": ["weight", "وزن"],
    "page_count": ["pages", "pagecount", "تعدادصفحه", "صفحات"],
}

REQUIRED_FIELDS = ("doc_code", "title")

MAX_REPORT_ROWS = 500


def _norm_header(value: object) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"[\s._\-:/\\()\[\]]+", "", text)
    return text


def suggest_mapping(headers: list[str]) -> dict[str, int]:
    """Map standard field names to 0-based column indexes."""
    normalised = [_norm_header(h) for h in headers]
    mapping: dict[str, int] = {}
    for field_name, aliases in COLUMN_ALIASES.items():
        alias_set = {_norm_header(a) for a in aliases}
        for index, header in enumerate(normalised):
            if not header or index in mapping.values():
                continue
            if header in alias_set:
                mapping[field_name] = index
                break
    return mapping


def detect_header_row(rows: list[list], max_scan: int = 15) -> int:
    """Return the 0-based index of the row that looks like the header."""
    best_index, best_score = 0, -1
    for index, row in enumerate(rows[:max_scan]):
        score = len(suggest_mapping([str(c) for c in row]))
        if score > best_score:
            best_index, best_score = index, score
    return best_index


def preview(file_path: str, sheet_name: str | None = None) -> dict:
    """Structure preview used by the mapping wizard."""
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheets = workbook.sheetnames
        worksheet = workbook[sheet_name] if sheet_name else workbook[sheets[0]]
        raw_rows = [
            [cell for cell in row]
            for row in worksheet.iter_rows(min_row=1, max_row=25, values_only=True)
        ]
    finally:
        workbook.close()
    header_index = detect_header_row(raw_rows)
    headers = [str(c).strip() if c is not None else "" for c in raw_rows[header_index]] if raw_rows else []
    return {
        "sheets": sheets,
        "active_sheet": sheet_name or (sheets[0] if sheets else None),
        "header_row": header_index + 1,  # 1-based for humans
        "headers": headers,
        "suggested_mapping": suggest_mapping(headers),
        "sample_rows": raw_rows[header_index + 1 : header_index + 6],
    }


@dataclass(slots=True)
class RowBox:
    row_number: int
    values: dict[str, object]


@dataclass(slots=True)
class ImportReport:
    status: str
    total: int = 0
    imported: int = 0
    updated: int = 0
    errors: int = 0
    warnings: int = 0
    entries: list[dict] = field(default_factory=list)
    batch_id: str | None = None

    def add(self, row: int, severity: Severity, message: str, column: str | None = None) -> None:
        if len(self.entries) >= MAX_REPORT_ROWS:
            return
        self.entries.append(
            {
                "row": row,
                "severity": severity.value,
                "column": column,
                "message": message,
            }
        )
        if severity is Severity.ERROR:
            self.errors += 1
        elif severity is Severity.WARNING:
            self.warnings += 1

    def as_dict(self) -> dict:
        return {
            "status": self.status,
            "batch_id": self.batch_id,
            "total_rows": self.total,
            "imported": self.imported,
            "updated": self.updated,
            "errors": self.errors,
            "warnings": self.warnings,
            "entries": self.entries,
        }


def _read_rows(file_path: str, sheet_name: str | None, header_row: int, mapping: dict[str, int]):
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        for row_index, row in enumerate(
            worksheet.iter_rows(min_row=header_row + 2, values_only=True), start=header_row + 2
        ):
            if row is None or all(cell is None or str(cell).strip() == "" for cell in row):
                continue
            values = {
                field_name: (row[index] if index < len(row) else None)
                for field_name, index in mapping.items()
            }
            yield RowBox(row_number=row_index, values=values)
    finally:
        workbook.close()


def _to_date(value: object) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, date):
        return value
    return parse_date(str(value))


def _to_float(value: object, default: float = 1.0) -> float:
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return default


def import_mdl(
    db: Session,
    *,
    project: Project,
    file_path: str,
    filename: str,
    mapping: dict[str, int],
    actor: User,
    sheet_name: str | None = None,
    header_row: int | None = None,
    auto_create_lookups: bool = True,
    dry_run: bool = False,
) -> ImportReport:
    """Validate and (unless ``dry_run``) atomically import the MDL."""
    cfg = merge_settings(project.settings)
    pattern = cfg["doc_code_pattern"]

    if "doc_code" not in mapping:
        report = ImportReport(status=MDLImportStatus.REJECTED.value)
        report.add(0, Severity.ERROR, "ستون کد مدرک در نگاشت مشخص نشده است.", "doc_code")
        return report

    sha = hashlib.sha256(Path(file_path).read_bytes()).hexdigest()
    header = header_row if header_row is not None else 0
    if header_row is None:
        with open(file_path, "rb") as handle:
            sample = io.BytesIO(handle.read())
        workbook = load_workbook(sample, read_only=True, data_only=True)
        try:
            worksheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
            rows = [[c for c in row] for row in worksheet.iter_rows(min_row=1, max_row=15, values_only=True)]
        finally:
            workbook.close()
        header = detect_header_row(rows)

    batch = MDLImportBatch(
        project_id=project.id,
        filename=filename,
        file_sha256=sha,
        status=MDLImportStatus.VALIDATING,
        mapping={k: v for k, v in mapping.items()},
        imported_by_id=actor.id,
    )
    db.add(batch)
    db.flush()

    report = ImportReport(status=MDLImportStatus.VALIDATING.value, batch_id=str(batch.id))
    seen_codes: dict[str, int] = {}
    pending: list[dict] = []

    disciplines = {
        d.code.strip().upper(): d
        for d in db.execute(select(Discipline).where(Discipline.project_id == project.id)).scalars()
    }
    units = {
        u.code.strip().upper(): u
        for u in db.execute(select(Unit).where(Unit.project_id == project.id)).scalars()
    }
    existing = {
        item.doc_code_norm: item
        for item in db.execute(select(MDLItem).where(MDLItem.project_id == project.id)).scalars()
    }

    for row in _read_rows(file_path, sheet_name, header, mapping):
        report.total += 1
        raw_code = row.values.get("doc_code")
        code_norm = normalize_code(str(raw_code) if raw_code is not None else "")
        if not code_norm:
            report.add(row.row_number, Severity.ERROR, "کد مدرک خالی است.", "doc_code")
            continue
        if not is_valid_code(code_norm, pattern):
            report.add(
                row.row_number, Severity.ERROR,
                f"کد «{raw_code}» با الگوی شماره‌گذاری پروژه سازگار نیست.", "doc_code",
            )
        if code_norm in seen_codes:
            report.add(
                row.row_number, Severity.ERROR,
                f"کد تکراری در فایل (سطر {seen_codes[code_norm]}).", "doc_code",
            )
            continue
        seen_codes[code_norm] = row.row_number

        title = str(row.values.get("title") or "").strip()
        if not title:
            report.add(row.row_number, Severity.WARNING, "عنوان مدرک خالی است.", "title")

        discipline_code = str(row.values.get("discipline") or "").strip().upper() or None
        unit_code = str(row.values.get("unit") or "").strip().upper() or None
        discipline = disciplines.get(discipline_code) if discipline_code else None
        unit = units.get(unit_code) if unit_code else None

        if discipline_code and discipline is None and not auto_create_lookups:
            report.add(row.row_number, Severity.WARNING,
                       f"دیسیپلین «{discipline_code}» در پروژه تعریف نشده است.", "discipline")

        rev_raw = row.values.get("current_rev")
        rev = normalize_rev(str(rev_raw) if rev_raw is not None else "")
        if rev_raw not in (None, "") and not rev:
            report.add(row.row_number, Severity.WARNING,
                       f"مقدار بازنگری «{rev_raw}» قابل تفسیر نیست.", "current_rev")

        due = _to_date(row.values.get("due_date"))
        if row.values.get("due_date") not in (None, "") and due is None:
            report.add(row.row_number, Severity.WARNING, "تاریخ موعد قابل تفسیر نیست.", "due_date")
        planned = _to_date(row.values.get("planned_issue_date"))

        pending.append(
            {
                "row": row.row_number,
                "doc_code": str(raw_code).strip(),
                "doc_code_norm": code_norm,
                "title": title or code_norm,
                "title_norm": normalize_title(title),
                "discipline_code": discipline_code,
                "unit_code": unit_code,
                "doc_type": str(row.values.get("doc_type") or "").strip() or None,
                "current_rev": rev or None,
                "rev_seq": rev_to_seq(rev, cfg["rev_scheme"]) if rev else 0,
                "status": str(row.values.get("status") or "").strip() or None,
                "weight": _to_float(row.values.get("weight"), 1.0),
                "page_count": int(_to_float(row.values.get("page_count"), 0)) or None,
                "due_date": due,
                "planned_issue_date": planned,
                "responsible_email": str(row.values.get("responsible") or "").strip() or None,
            }
        )

    # ------------------------------------------------------------- commit
    if report.errors and not dry_run:
        report.status = MDLImportStatus.REJECTED.value
        batch.status = MDLImportStatus.REJECTED
        batch.row_count = report.total
        batch.error_count = report.errors
        batch.warning_count = report.warnings
        batch.report = report.entries
        db.flush()
        return report

    if dry_run:
        report.status = "validated"
        report.imported = len(pending)
        db.rollback()
        return report

    # ensure lookups
    for row in pending:
        for kind, code, registry in (
            ("discipline", row["discipline_code"], disciplines),
            ("unit", row["unit_code"], units),
        ):
            if not code or code in registry:
                continue
            if not auto_create_lookups:
                continue
            if kind == "discipline":
                new_lookup = Discipline(project_id=project.id, code=code, name=code)
            else:
                new_lookup = Unit(project_id=project.id, code=code, name=code)
            db.add(new_lookup)
            db.flush()
            registry[code] = new_lookup

    users_by_email = {
        u.email.lower(): u for u in db.execute(select(User)).scalars() if u.email
    }

    for row in pending:
        discipline = disciplines.get(row["discipline_code"] or "")
        unit = units.get(row["unit_code"] or "")
        responsible = None
        if row["responsible_email"]:
            responsible = users_by_email.get(row["responsible_email"].lower())

        item = existing.get(row["doc_code_norm"])
        if item is None:
            item = MDLItem(
                project_id=project.id,
                import_batch_id=batch.id,
                doc_code=row["doc_code"],
                doc_code_norm=row["doc_code_norm"],
                title=row["title"],
                title_norm=row["title_norm"],
                discipline_id=discipline.id if discipline else None,
                unit_id=unit.id if unit else None,
                doc_type=row["doc_type"],
                current_rev=row["current_rev"],
                rev_seq=row["rev_seq"],
                status=row["status"],
                responsible_id=responsible.id if responsible else None,
                planned_issue_date=row["planned_issue_date"],
                due_date=row["due_date"],
                weight=row["weight"],
                page_count=row["page_count"],
                source_row=row["row"],
                aliases=[],
                attributes={},
            )
            db.add(item)
            existing[row["doc_code_norm"]] = item
            report.imported += 1
        else:
            item.title = row["title"] or item.title
            item.title_norm = row["title_norm"] or item.title_norm
            item.discipline_id = discipline.id if discipline else item.discipline_id
            item.unit_id = unit.id if unit else item.unit_id
            item.doc_type = row["doc_type"] or item.doc_type
            item.current_rev = row["current_rev"] or item.current_rev
            item.rev_seq = row["rev_seq"] if row["current_rev"] else item.rev_seq
            item.status = row["status"] or item.status
            item.weight = row["weight"]
            item.due_date = row["due_date"] or item.due_date
            item.planned_issue_date = row["planned_issue_date"] or item.planned_issue_date
            item.page_count = row["page_count"] or item.page_count
            item.import_batch_id = batch.id
            item.source_row = row["row"]
            item.is_active = True
            report.updated += 1

    batch.status = MDLImportStatus.COMPLETED
    batch.row_count = report.total
    batch.imported_count = report.imported
    batch.error_count = report.errors
    batch.warning_count = report.warnings
    batch.report = report.entries
    report.status = MDLImportStatus.COMPLETED.value

    audit.record(
        db,
        action="mdl.import",
        entity_type="mdl_import_batch",
        entity_id=batch.id,
        actor=actor,
        project_id=project.id,
        after={
            "filename": filename,
            "imported": report.imported,
            "updated": report.updated,
            "warnings": report.warnings,
        },
    )
    db.flush()
    return report


def export_mdl_bytes(items: list[MDLItem]) -> bytes:
    """Excel export in a format that can be re-imported unchanged."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "MDL"
    headers = [
        "Document No", "Title", "Discipline", "Unit", "Type", "Current Rev",
        "Status", "Pages", "Weight", "Due Date",
    ]
    worksheet.append(headers)
    for item in items:
        worksheet.append(
            [
                item.doc_code,
                item.title,
                item.discipline.code if item.discipline else "",
                item.unit.code if item.unit else "",
                item.doc_type or "",
                item.current_rev or "",
                item.status or "",
                item.page_count or "",
                item.weight,
                item.due_date.isoformat() if item.due_date else "",
            ]
        )
    worksheet.freeze_panes = "A2"
    for column, width in zip("ABCDEFGHIJ", (24, 48, 12, 10, 12, 12, 12, 8, 8, 14)):
        worksheet.column_dimensions[column].width = width
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
