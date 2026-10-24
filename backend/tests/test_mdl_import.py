"""MDL Excel import: mapping detection, validation report, atomicity."""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook
from sqlalchemy.orm import Session

from app.models import MDLItem, Project, User
from app.services import mdl_import
from app.services.project_settings import merge_settings


def build_xlsx(rows: list[list], headers: list[str] | None = None) -> str:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "MDL"
    sheet.append(headers or ["Document No", "Title", "Discipline", "Unit", "Type", "Rev", "Status"])
    for row in rows:
        sheet.append(row)
    path = f"/tmp/mdl-test-{abs(hash(str(rows)))}.xlsx"
    workbook.save(path)
    return path


@pytest.fixture()
def actor(db: Session) -> User:
    user = db.query(User).first()
    if user is None:
        user = User(email="importer@edms-test.com", full_name="Importer", hashed_password="x")
        db.add(user)
        db.commit()
    return user


def test_preview_detects_headers(db, actor):
    path = build_xlsx([["TST-CIV-DR-0101", "Plot Plan", "CIV", "U-100", "DWG", "A", "Draft"]])
    preview = mdl_import.preview(path)
    assert preview["sheets"] == ["MDL"]
    assert preview["header_row"] == 1
    mapping = preview["suggested_mapping"]
    assert mapping["doc_code"] == 0
    assert mapping["title"] == 1
    assert mapping["discipline"] == 2


def test_preview_persian_headers(db, actor):
    path = build_xlsx(
        [["TST-CIV-DR-0102", "نقشه آزمون", "CIV", "U-100", "DWG", "B", "Draft"]],
        headers=["کد مدرک", "عنوان", "دیسیپلین", "یونیت", "نوع", "بازنگری", "وضعیت"],
    )
    mapping = mdl_import.preview(path)["suggested_mapping"]
    assert mapping["doc_code"] == 0
    assert mapping["current_rev"] == 5


def test_import_success(db: Session, project, actor):
    path = build_xlsx(
        [
            ["TST-CIV-DR-0101", "Plot Plan", "CIV", "U-100", "DWG", "A", "Draft"],
            ["TST-CIV-DR-0102", "Site Grading", "CIV", "U-100", "DWG", "B", "Draft"],
        ]
    )
    preview = mdl_import.preview(path)
    report = mdl_import.import_mdl(
        db,
        project=project,
        file_path=path,
        filename="mdl.xlsx",
        mapping=preview["suggested_mapping"],
        actor=actor,
        header_row=preview["header_row"] - 1,
    )
    db.commit()
    assert report.status == "completed"
    assert report.imported == 2
    assert report.errors == 0
    assert db.query(MDLItem).filter_by(project_id=project.id, doc_code_norm="TST-CIV-DR-0101").count() == 1


def test_import_rejects_invalid_code_pattern_and_writes_nothing(db: Session, project, actor):
    path = build_xlsx([["BAD CODE!!", "Invalid", "CIV", "U-100", "DWG", "A", "Draft"]])
    preview = mdl_import.preview(path)
    before = db.query(MDLItem).filter_by(project_id=project.id).count()
    report = mdl_import.import_mdl(
        db, project=project, file_path=path, filename="bad.xlsx",
        mapping=preview["suggested_mapping"], actor=actor, header_row=preview["header_row"] - 1,
    )
    assert report.status == "rejected"
    assert report.errors >= 1
    assert db.query(MDLItem).filter_by(project_id=project.id).count() == before


def test_import_detects_duplicate_rows(db: Session, project, actor):
    path = build_xlsx(
        [
            ["TST-CIV-DR-0201", "Dup", "CIV", "U-100", "DWG", "A", "Draft"],
            ["TST-CIV-DR-0201", "Dup again", "CIV", "U-100", "DWG", "A", "Draft"],
        ]
    )
    preview = mdl_import.preview(path)
    report = mdl_import.import_mdl(
        db, project=project, file_path=path, filename="dup.xlsx",
        mapping=preview["suggested_mapping"], actor=actor, header_row=preview["header_row"] - 1,
    )
    assert report.status == "rejected"
    assert any("تکراری" in entry["message"] for entry in report.entries)


def test_import_updates_existing_row(db: Session, project, actor):
    path = build_xlsx([["TST-CIV-DR-001", "Plot Plan (Rev C)", "CIV", "U-100", "DWG", "C", "Draft"]])
    preview = mdl_import.preview(path)
    report = mdl_import.import_mdl(
        db, project=project, file_path=path, filename="update.xlsx",
        mapping=preview["suggested_mapping"], actor=actor, header_row=preview["header_row"] - 1,
    )
    db.commit()
    assert report.updated == 1
    item = db.query(MDLItem).filter_by(doc_code_norm="TST-CIV-DR-001").one()
    assert item.current_rev == "C"
    assert "Rev C" in item.title


def test_dry_run_does_not_persist(db: Session, project, actor):
    path = build_xlsx([["TST-ELE-DR-0401", "SLD", "ELE", "U-300", "DWG", "A", "Draft"]])
    preview = mdl_import.preview(path)
    before = db.query(MDLItem).filter_by(project_id=project.id).count()
    report = mdl_import.import_mdl(
        db, project=project, file_path=path, filename="dry.xlsx",
        mapping=preview["suggested_mapping"], actor=actor,
        header_row=preview["header_row"] - 1, dry_run=True,
    )
    assert report.status == "validated"
    assert db.query(MDLItem).filter_by(project_id=project.id).count() == before


def test_missing_code_column_is_rejected(db: Session, project, actor):
    path = build_xlsx([["Title only"]], headers=["Title"])
    report = mdl_import.import_mdl(
        db, project=project, file_path=path, filename="nocol.xlsx",
        mapping={"title": 0}, actor=actor,
    )
    assert report.status == "rejected"


def test_export_roundtrip(db: Session, project):
    items = db.query(MDLItem).filter_by(project_id=project.id).all()
    payload = mdl_import.export_mdl_bytes(items)
    assert payload[:2] == b"PK"
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(payload))
    sheet = workbook.active
    assert sheet.max_row == len(items) + 1
